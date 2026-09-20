"""Turn any vendor file into *anchored text*: plain text where every line
carries a location tag (`[Quotation!F6]`, `[page 2]`, `[para 7]`,
`[image line 12]`). The extractor must quote these locations, which is what
makes the evidence drawer possible.

Excel, PDF, Word and email are read deterministically. Photos (and PDFs with no
text layer) are transcribed by the vision model; the transcription is labelled
as such and its legibility score is carried through to the UI.
"""
from __future__ import annotations

import email
import email.policy
import io
import mimetypes
import re

import pymupdf
from docx import Document
from openpyxl import load_workbook

from . import llm
from .models import ImageTranscriptionAI

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif"}


def guess_content_type(name: str) -> str:
    ct, _ = mimetypes.guess_type(name)
    if name.lower().endswith(".eml"):
        return "message/rfc822"
    return ct or "application/octet-stream"


def kind_from_name(name: str) -> str:
    n = name.lower()
    if n.endswith((".xlsx", ".xlsm", ".xls")):
        return "xlsx"
    if n.endswith(".pdf"):
        return "pdf"
    if n.endswith(".docx"):
        return "docx"
    if any(n.endswith(e) for e in IMAGE_EXT):
        return "image"
    if n.endswith((".eml", ".txt", ".msg")):
        return "email"
    if n.endswith(".csv"):
        return "csv"
    return "unknown"


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def xlsx_to_text(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Sheet: {ws.title} ===")
        for row in ws.iter_rows():
            cells = [(c.coordinate, c.value) for c in row if c.value not in (None, "")]
            if not cells:
                continue
            out.append(" | ".join(f"[{ws.title}!{coord}] {_fmt_cell(v)}" for coord, v in cells))
    return "\n".join(out)


def csv_to_text(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    return "\n".join(f"[row {i}] {line}" for i, line in enumerate(text.splitlines(), start=1) if line.strip())


def _pdf_page_rows(page, y_tol: float = 3.0) -> list[str]:
    """Rebuild visual rows from positioned words so table cells on one line stay on one line.

    pymupdf's plain text mode emits each table cell as its own line; grouping words by their
    vertical position gives the model (and the evidence check) the row a human would read.
    """
    words = page.get_text("words")  # x0, y0, x1, y1, word, block, line, wordno
    if not words:
        return []
    words.sort(key=lambda w: (round(w[1] / y_tol), w[0]))
    rows: list[list] = []
    for w in words:
        if rows and abs(rows[-1][0][1] - w[1]) <= y_tol:
            rows[-1].append(w)
        else:
            rows.append([w])
    out = []
    for r in rows:
        r.sort(key=lambda w: w[0])
        parts = []
        prev_x1 = None
        for w in r:
            gap = (w[0] - prev_x1) if prev_x1 is not None else 0
            parts.append(("  " if gap > 12 else " ") + w[4] if parts else w[4])
            prev_x1 = w[2]
        out.append("".join(parts).strip())
    return out


def pdf_to_text(data: bytes) -> tuple[str, bool]:
    """Returns (text, has_text_layer)."""
    doc = pymupdf.open(stream=data, filetype="pdf")
    out = []
    total = 0
    for i, page in enumerate(doc, start=1):
        rows = _pdf_page_rows(page)
        total += sum(len(r) for r in rows)
        out.append(f"=== [page {i}] ===")
        for line in rows:
            if line.strip():
                out.append(f"[page {i}] {line}")
    doc.close()
    return "\n".join(out), total > 40


def pdf_page_images(data: bytes, max_pages: int = 4, zoom: float = 1.6) -> list[bytes]:
    doc = pymupdf.open(stream=data, filetype="pdf")
    imgs = []
    for page in list(doc)[:max_pages]:
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
        imgs.append(pix.tobytes("png"))
    doc.close()
    return imgs


def docx_to_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    out = []
    p_i = 0
    for p in d.paragraphs:
        if p.text.strip():
            p_i += 1
            prefix = "heading" if p.style.name.lower().startswith("heading") else "para"
            out.append(f"[{prefix} {p_i}] {p.text.strip()}")
    for t_i, t in enumerate(d.tables, start=1):
        for r_i, row in enumerate(t.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            out.append(f"[table {t_i} row {r_i}] " + " | ".join(cells))
    return "\n".join(out)


def email_to_text(data: bytes) -> str:
    raw = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
    headers: list[str] = []
    body_text = raw
    head, sep, rest = raw.partition("\n\n")
    if sep and re.match(r"^(From|To|Subject|Date|Received|Return-Path|Message-ID):", head, re.I | re.M):
        try:
            msg = email.message_from_string(head + "\n\n", policy=email.policy.default)
            headers = [f"[header] {k}: {msg[k]}" for k in ("From", "To", "Date", "Subject") if msg[k]]
        except Exception:
            headers = []
        # Take the body straight from the raw text so non-ASCII (₹) survives missing charset declarations.
        body_text = rest
        if "Content-Type:" in head and "multipart" in head.lower():
            try:
                msg = email.message_from_string(raw, policy=email.policy.default)
                body = msg.get_body(preferencelist=("plain", "html"))
                body_text = body.get_content() if body else rest
            except Exception:
                body_text = rest
    lines = [f"[line {i}] {ln.rstrip()}" for i, ln in enumerate(body_text.splitlines(), start=1) if ln.strip()]
    return "\n".join(headers + lines)


def image_to_text(data: bytes, media_type: str, log: list | None = None) -> dict:
    system = (
        "You transcribe photographed business documents for a procurement team. Transcribe faithfully, line by line, "
        "preserving numbers exactly as printed. Use ' | ' to separate table columns. Where a digit or word is not "
        "legible write '?'. Where a row is cut off by the frame write '[cut off]'. Include small print and footnotes. "
        "Never guess or 'correct' a number."
    )
    res = llm.structured(
        purpose="transcribe_image",
        system=system,
        content=[llm.image_block(data, media_type), llm.text_block("Transcribe this document.")],
        schema=ImageTranscriptionAI,
        max_tokens=6000,
        log=log,
    )
    lines = [f"[image line {i}] {ln.rstrip()}" for i, ln in enumerate(res.transcription.splitlines(), start=1) if ln.strip()]
    return {"text": "\n".join(lines), "method": "vision-transcription", "legibility": res.legibility, "caveats": res.caveats}


def file_to_text(name: str, data: bytes, kind: str | None = None, log: list | None = None) -> dict:
    """Returns {text, method, legibility?, caveats?}."""
    kind = kind or kind_from_name(name)
    if kind == "xlsx":
        return {"text": xlsx_to_text(data), "method": "openpyxl cell dump"}
    if kind == "csv":
        return {"text": csv_to_text(data), "method": "csv rows"}
    if kind == "pdf":
        text, has_layer = pdf_to_text(data)
        if has_layer:
            return {"text": text, "method": "pymupdf text layer"}
        parts = []
        caveats = []
        leg = []
        for i, png in enumerate(pdf_page_images(data), start=1):
            t = image_to_text(png, "image/png", log=log)
            parts.append(f"=== [page {i}] (vision transcription) ===\n" + t["text"].replace("[image line", f"[page {i} line"))
            caveats.append(t.get("caveats", ""))
            leg.append(t.get("legibility", 0))
        return {"text": "\n".join(parts), "method": "vision-transcription (scanned PDF)", "legibility": min(leg) if leg else 0, "caveats": " ".join(c for c in caveats if c)}
    if kind == "docx":
        return {"text": docx_to_text(data), "method": "python-docx paragraphs and tables"}
    if kind == "image":
        media = guess_content_type(name)
        if media not in ("image/jpeg", "image/png", "image/webp", "image/gif"):
            media = "image/jpeg"
        return image_to_text(data, media, log=log)
    if kind == "email":
        return {"text": email_to_text(data), "method": "email body"}
    # Unknown: try as text
    try:
        txt = data.decode("utf-8")
        return {"text": "\n".join(f"[line {i}] {l}" for i, l in enumerate(txt.splitlines(), 1) if l.strip()), "method": "plain text"}
    except UnicodeDecodeError:
        raise ValueError(f"Unsupported file type: {name}")


def normalize_ws(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip().lower()
