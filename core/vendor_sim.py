"""Synthetic vendor replies, generated from the *actual* drafted RFx.

Vendor personalities and their ugly formats are fixed; the line items, sizes
and quantities follow whatever the AI drafted. Nothing here is read back by
the extractor: the extractor only ever sees the produced files, exactly like
a real inbox.

Pricing is derived from a simple board-cost model (rupees per kg of board,
scaled by nominal carton weight, plus print surcharges), perturbed per vendor,
so that cheapest-per-line splits are genuinely interesting.
"""
from __future__ import annotations

import io
import math
import random
import uuid
from datetime import datetime, timedelta, timezone

import numpy as np
import pymupdf
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from . import storage

EXAMPLE_BRIEF = (
    "We are a packaged snacks manufacturer with a plant in Chakan (Pune). For the next financial year we need "
    "about 30 SKUs of corrugated packaging: outer shippers for chips and namkeen (mostly 3-ply, some 5-ply for "
    "export and heavier loads), a few die-cut display trays, and some 7-ply master cartons for palletised export. "
    "Total spend last year was around Rs 3.8 crore. Deliveries weekly to Chakan, prices delivered and exclusive of GST, "
    "60-day validity, 45-day payment. Vendors must have ISO 9001 and be able to provide BCT test reports; FSC or "
    "recycled-content declaration preferred. Print is mostly 1-2 colour flexo with our brand marks."
)

VENDORS = [
    {"vendor_id": "v1", "name": "Sri Balaji Packaging Pvt Ltd", "city": "Pune", "email": "quotes@sribalajipack.in", "format": "xlsx"},
    {"vendor_id": "v2", "name": "Kraftline Industries", "city": "Nashik", "email": "sales@kraftline.co.in", "format": "pdf"},
    {"vendor_id": "v3", "name": "PakAsia Global Pte Ltd", "city": "Singapore", "email": "rfq@pakasia.sg", "format": "docx"},
    {"vendor_id": "v4", "name": "Meghna Corrugators", "city": "Ahmedabad", "email": "meghnacorru@gmail.com", "format": "image"},
    {"vendor_id": "v5", "name": "Ganesh Board Mills", "city": "Indore", "email": "ganeshboard.indore@gmail.com", "format": "email"},
]

FILL_HEAD = PatternFill("solid", fgColor="1F3A5F")
FILL_ALT = PatternFill("solid", fgColor="EEF3F9")
THIN = Side(style="thin", color="B8C4D6")


def cover_email(rfx: dict, vendor: dict) -> str:
    due = (datetime.now(timezone.utc) + timedelta(days=9)).strftime("%d %b %Y")
    return (
        f"Dear {vendor['name']} team,\n\n"
        f"Please find attached our request for quotation: {rfx['title']}.\n\n"
        f"Scope: {len(rfx['line_items'])} corrugated packaging line items (specifications attached).\n"
        f"Price basis: {rfx['terms']['price_basis']}.\nPayment: {rfx['terms']['payment_terms']}.\n"
        f"Delivery: {rfx['terms']['delivery_location']}; {rfx['terms']['delivery_schedule']}.\n"
        f"Validity required: {rfx['terms']['validity_days']} days.\n\n"
        f"Please also complete the {len(rfx['questionnaire'])}-question quality questionnaire and attach certificates where requested.\n"
        f"Quotes are due by {due}. You may reply in any format; we will normalise them on our side.\n\n"
        "Regards,\nCategory Sourcing"
    )


def empty_vendor(v: dict) -> dict:
    return {**v, "files": [], "texts": {}, "extraction": None, "status": "awaiting", "error": None, "received_at": None}


# ---------------------------------------------------------------------------
# Pricing model
# ---------------------------------------------------------------------------

def _base_price_inr(li: dict, rng: random.Random) -> float:
    """Fair per-piece price in INR from a board-cost model. Deterministic given rng state."""
    board_rate_per_kg = {"3-ply": 46.0, "5-ply": 50.0, "7-ply": 54.0}.get(li["board"], 48.0)
    weight_kg = li["nominal_weight_g"] / 1000
    print_surcharge = 0.0
    p = (li.get("print") or "").lower()
    if "2" in p and "colour" in p or "2 color" in p:
        print_surcharge = 0.9
    elif "colour" in p or "color" in p or "print" in p and "unprint" not in p:
        print_surcharge = 0.5
    conv = 1.18  # conversion, glue, margin
    return max(3.0, weight_kg * board_rate_per_kg * conv + print_surcharge)


def _vendor_prices(rfx: dict, seed: int) -> dict[str, dict[int, float]]:
    """Per-vendor INR-per-piece prices (before their own format quirks)."""
    rng = random.Random(seed)
    out: dict[str, dict[int, float]] = {v["vendor_id"]: {} for v in VENDORS}
    factors = {"v1": 1.00, "v2": 0.97, "v3": 1.06, "v4": 0.95, "v5": 0.93}
    for li in rfx["line_items"]:
        base = _base_price_inr(li, rng)
        for vid, f in factors.items():
            noise = rng.uniform(-0.06, 0.06)
            # each vendor is strong in a different board grade
            grade_bias = {
                "v1": {"5-ply": -0.03},
                "v2": {"3-ply": -0.04, "7-ply": 0.08},
                "v3": {"7-ply": -0.10, "3-ply": 0.05},
                "v4": {"5-ply": 0.04},
                "v5": {"3-ply": -0.02},
            }[vid].get(li["board"], 0.0)
            out[vid][li["line_no"]] = round(base * (f + noise + grade_bias), 2)
    return out


def _round_up(x: float, step: float) -> float:
    return math.ceil(x / step) * step


def _text_answer(question: str) -> str:
    """A plausible free-text questionnaire answer keyed on what the question asks about."""
    q = question.lower()
    if "capacity" in q:
        return "3 corrugation lines, 1,800 MT/month installed; ~72% utilised"
    if "reference" in q or "experience" in q or "customer" in q:
        return "Parle Agro (Mr. S. Kulkarni, 98220 11223), Haldiram's Nagpur (Ms. R. Deshmukh, 98230 44556)"
    if "print" in q or "flexo" in q or "colour" in q:
        return "4-colour flexo folder-gluer; plate change 40 min; spectrophotometer checks per shift"
    if "capa" in q or "complaint" in q or "corrective" in q:
        return "Documented CAPA under ISO 9001; average closure 7 working days"
    if "traceab" in q or "lot" in q or "batch" in q:
        return "Yes, batch code and date printed on every carton flap"
    if "test" in q or "bct" in q or "ect" in q or "quality control" in q or "inspection" in q:
        return "In-house lab with BCT, ECT and Cobb testers; reports issued per lot; AQL 1.0 sampling"
    return "Yes, details available on request"


# ---------------------------------------------------------------------------
# V1: the beautiful Excel that ignores the template
# ---------------------------------------------------------------------------

def build_v1_xlsx(rfx: dict, prices: dict[int, float], rng: random.Random) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.merge_cells("A1:G1")
    ws["A1"] = "SRI BALAJI PACKAGING PVT LTD - COMMERCIAL OFFER"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = FILL_HEAD
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Ref: SBP/Q/2026-27/0418"
    ws["E2"] = f"Date: {datetime.now().strftime('%d-%b-%Y')}"
    ws["A3"] = f"Subject: Offer against your enquiry - {rfx['title']}"
    headers = ["Item Code", "Product", "Size (L x W x H) mm", "Ply / Flute", "Board GSM", "Rate per 100 Nos (Rs.)", "MOQ (Nos)"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    alt_lines = set(rng.sample([li["line_no"] for li in rfx["line_items"] if li["board"] == "5-ply"] or [3, 12], k=2))
    r = 6
    for li in rfx["line_items"]:
        per100 = round(prices[li["line_no"]] * 100, 0)
        product = li["description"]
        flute = li["flute"]
        gsm = li["gsm"]
        if li["line_no"] in alt_lines:
            product += " (ALT: offered in BC flute instead of specified flute, same GSM)"
            flute = "BC"
        row = [f"SBP-{1000 + li['line_no'] * 7}", product, f"{li['length_mm']} x {li['width_mm']} x {li['height_mm']}", f"{li['board']} / {flute}", gsm, per100, 500 if li["annual_qty"] > 20000 else 1000]
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if (r % 2) == 0:
                cell.fill = FILL_ALT
            if c == 6:
                cell.number_format = "#,##0.00"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Notes:").font = Font(bold=True)
    notes = [
        "1. Rates are per 100 pieces, delivered to your Chakan plant, GST extra @ 12%.",
        "2. Special discount of 3% on invoice value if the annual PO is released before 30-Sep-2026.",
        "3. Payment: 45 days from date of invoice. Offer validity: 60 days.",
        "4. Items marked ALT are offered in BC flute; please confirm acceptance.",
        "5. Kraft liner sourced from FSC-CoC certified mills; certificate attached.",
    ]
    for n in notes:
        r += 1
        ws.cell(row=r, column=1, value=n)
    widths = [14, 58, 22, 16, 11, 22, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Compliance")
    ws2["A1"] = "Question"
    ws2["B1"] = "Response"
    ws2["C1"] = "Remarks"
    for c in "ABC":
        ws2[f"{c}1"].font = Font(bold=True)
    qs = rfx["questionnaire"]
    skipped = set(rng.sample([q["q_id"] for q in qs], k=min(2, max(0, len(qs) - 6))))
    for i, q in enumerate(qs, start=2):
        ws2.cell(row=i, column=1, value=f"{q['q_id']}: {q['text']}")
        if q["q_id"] in skipped:
            continue
        if q["answer_type"] == "yes_no":
            ws2.cell(row=i, column=2, value="Yes")
        elif q["answer_type"] == "document":
            ws2.cell(row=i, column=2, value="Attached")
            ws2.cell(row=i, column=3, value="See ISO 9001:2015 certificate PDF")
        else:
            ws2.cell(row=i, column=2, value=_text_answer(q["text"]))
    ws2.column_dimensions["A"].width = 80
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 50
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# V2: PDF on letterhead, 27 of 30 lines, discount in a footnote
# ---------------------------------------------------------------------------

def build_v2_pdf(rfx: dict, prices: dict[int, float], rng: random.Random) -> tuple[bytes, list[int]]:
    skipped = sorted(rng.sample([li["line_no"] for li in rfx["line_items"]], k=3))
    doc = pymupdf.open()
    W, H = 595, 842
    page = doc.new_page(width=W, height=H)

    def letterhead(pg):
        pg.draw_rect(pymupdf.Rect(0, 0, W, 70), color=None, fill=(0.55, 0.25, 0.1))
        pg.insert_text((40, 38), "KRAFTLINE INDUSTRIES", fontsize=20, fontname="hebo", color=(1, 1, 1))
        pg.insert_text((40, 56), "Corrugated Boxes  |  Plot 14, MIDC Satpur, Nashik 422007  |  GSTIN 27AAACK1234F1Z5", fontsize=8, fontname="helv", color=(1, 1, 1))

    letterhead(page)
    y = 95
    page.insert_text((40, y), f"Ref: KI/QTN/{datetime.now().strftime('%y%m')}/117            Date: {datetime.now().strftime('%d.%m.%Y')}", fontsize=9, fontname="helv")
    y += 18
    page.insert_text((40, y), "To: Category Sourcing", fontsize=9, fontname="helv")
    y += 18
    page.insert_text((40, y), f"Sub: Quotation for {rfx['title']}", fontsize=10, fontname="hebo")
    y += 16
    page.insert_textbox(
        pymupdf.Rect(40, y, W - 40, y + 40),
        "Dear Sir/Madam, we thank you for your enquiry and are pleased to quote as under. Rates are in Indian Rupees per piece*, ex-works Nashik**.",
        fontsize=9,
        fontname="helv",
    )
    y += 44
    cols = [40, 70, 330, 420, 480, 555]
    heads = ["Sr", "Description", "Size (mm)", "Ply", "Rate (Rs.)"]

    def header_row(pg, yy):
        pg.draw_rect(pymupdf.Rect(40, yy - 11, 555, yy + 4), color=None, fill=(0.93, 0.9, 0.86))
        for i, h in enumerate(heads):
            pg.insert_text((cols[i] + 3, yy), h, fontsize=8.5, fontname="hebo")
        return yy + 14

    y = header_row(page, y)
    sr = 0
    for li in rfx["line_items"]:
        if li["line_no"] in skipped:
            continue
        sr += 1
        if y > H - 120:
            page = doc.new_page(width=W, height=H)
            letterhead(page)
            y = 95
            y = header_row(page, y)
        desc = f"{li['description']} - {li['flute']} flute, {li['gsm']} gsm"
        if len(desc) > 62:
            desc = desc[:60] + ".."
        vals = [str(sr), desc, f"{li['length_mm']}x{li['width_mm']}x{li['height_mm']}", li["board"], f"{prices[li['line_no']]:.2f}"]
        for i, v in enumerate(vals):
            page.insert_text((cols[i] + 3, y), v, fontsize=8.5, fontname="helv")
        page.draw_line((40, y + 4), (555, y + 4), color=(0.8, 0.8, 0.8), width=0.3)
        y += 14
    y += 10
    page.insert_text((40, y), f"Items not quoted: your line items {', '.join(str(s) for s in skipped)} are outside our current die range.", fontsize=8.5, fontname="helv")
    y += 22
    page.insert_text((40, y), "Commercial terms", fontsize=10, fontname="hebo")
    y += 14
    terms = [
        "Payment: 30 days from receipt of material.",
        "Validity: 45 days from the date of this quotation.",
        "Lead time: 10-12 working days from approved artwork.",
        "GST: extra as applicable (currently 12% on HSN 4819).",
    ]
    for t in terms:
        page.insert_text((40, y), t, fontsize=8.5, fontname="helv")
        y += 12
    y += 8
    page.insert_text((40, y), "Compliance statement", fontsize=10, fontname="hebo")
    y += 14
    qs = rfx["questionnaire"]
    answers = []
    for q in qs:
        if q["answer_type"] == "yes_no":
            answers.append(f"{q['q_id']}: Yes.")
        elif q["answer_type"] == "document":
            answers.append(f"{q['q_id']}: FSC chain-of-custody certificate enclosed; ISO 9001 certification is under renewal, copy to follow.")
        else:
            answers.append(f"{q['q_id']}: In-house testing for BCT, ECT and Cobb; third-party reports from SGS on request.")
    text = " ".join(answers)
    rect = pymupdf.Rect(40, y - 8, W - 40, y + 130)
    page.insert_textbox(rect, text, fontsize=8, fontname="helv")
    # footnotes, small, at the bottom of the last page
    fy = H - 48
    page.draw_line((40, fy - 8), (300, fy - 8), color=(0.5, 0.5, 0.5), width=0.4)
    page.insert_text((40, fy), "* A discount of 5% applies on all above rates for an annual commitment exceeding Rs. 25 lakh.", fontsize=6.5, fontname="helv", color=(0.3, 0.3, 0.3))
    page.insert_text((40, fy + 9), "** Freight extra at actuals from Nashik to your plant; transit insurance to buyer's account.", fontsize=6.5, fontname="helv", color=(0.3, 0.3, 0.3))
    page.insert_text((40, fy + 18), "E&OE. This quotation is subject to Kraftline standard terms of sale.", fontsize=6.5, fontname="helv", color=(0.3, 0.3, 0.3))
    out = doc.tobytes(deflate=True)
    doc.close()
    return out, skipped


# ---------------------------------------------------------------------------
# V3: Word doc, commercials in paragraphs, USD per 1000 pcs
# ---------------------------------------------------------------------------

def build_v3_docx(rfx: dict, prices: dict[int, float], usd_rate: float, rng: random.Random) -> bytes:
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    d.add_heading("PakAsia Global Pte Ltd", level=1)
    d.add_paragraph("71 Robinson Road, #14-01, Singapore 068895  |  Corrugated & Folding Carton Trading  |  UEN 201912345K")
    d.add_paragraph(f"Date: {datetime.now().strftime('%d %B %Y')}    Our ref: PAG-IN-{datetime.now().strftime('%y')}-0362")
    d.add_heading(f"Commercial proposal: {rfx['title']}", level=2)
    d.add_paragraph(
        "Thank you for inviting PakAsia Global to participate. We propose to supply from our partner mills in Gujarat and "
        "Vietnam, consolidated and shipped to Nhava Sheva. All prices below are quoted in United States Dollars per one "
        "thousand pieces, CIF Nhava Sheva, and are exclusive of Indian GST, import duties and inland freight to Chakan."
    )
    five = [li for li in rfx["line_items"] if li["board"] == "5-ply"]
    three = [li for li in rfx["line_items"] if li["board"] == "3-ply"]
    seven = [li for li in rfx["line_items"] if li["board"] not in ("3-ply", "5-ply")]

    def usd_per_1000(line_no: int) -> float:
        return round(prices[line_no] * 1000 / usd_rate, 1)

    d.add_heading("Five-ply and heavy-duty items", level=3)
    if five or seven:
        t = d.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        for i, h in enumerate(["Your line", "Item", "Dimensions (mm)", "USD / 1,000 pcs"]):
            hdr[i].text = h
        for li in five + seven:
            row = t.add_row().cells
            row[0].text = str(li["line_no"])
            row[1].text = f"{li['description']} ({li['board']}, {li['flute']} flute)"
            row[2].text = f"{li['length_mm']} x {li['width_mm']} x {li['height_mm']}"
            row[3].text = f"{usd_per_1000(li['line_no']):,.1f}"
    d.add_heading("Three-ply items", level=3)
    if three:
        # A flat paragraph price with exceptions: deliberately awkward to parse.
        exceptions = three[:2]
        rest = three[2:]
        flat = round(sum(usd_per_1000(li["line_no"]) for li in rest) / max(1, len(rest)), 1) if rest else usd_per_1000(three[0]["line_no"])
        rest_ids = ", ".join(str(li["line_no"]) for li in rest)
        p = d.add_paragraph()
        p.add_run(
            f"For the three-ply shippers we can offer a single blended rate of USD {flat:,.1f} per thousand pieces covering your lines "
            f"{rest_ids}, provided the lines are ordered together in the stated annual volumes. "
        )
        for li in exceptions:
            p.add_run(f"Line {li['line_no']} ({li['description']}) is priced separately at USD {usd_per_1000(li['line_no']):,.1f} per thousand owing to its print specification. ")
    d.add_heading("Terms", level=3)
    d.add_paragraph(
        "Payment is 30% advance with order and 70% against copy of shipping documents. Our proposal remains valid for 30 days. "
        "Lead time is 6 to 7 weeks from artwork approval to arrival at Nhava Sheva. Minimum order per line is 5,000 pieces. "
        "Should the total annual award exceed USD 250,000 we would extend a 2.5% volume rebate settled quarterly."
    )
    d.add_heading("Quality and compliance", level=3)
    qs = rfx["questionnaire"]
    sentences = []
    for q in qs:
        if q["answer_type"] == "document":
            sentences.append(f"Regarding {q['q_id']}, our partner mills hold ISO 9001:2015; certified copies are available on request.")
        elif q["answer_type"] == "yes_no":
            sentences.append(f"To {q['q_id']} our answer is yes." if rng.random() > 0.2 else f"On {q['q_id']} we are unable to commit at this stage.")
        else:
            sentences.append(f"For {q['q_id']}: testing is performed at the mill laboratory with BCT and ECT equipment; we supply COAs with every shipment.")
    d.add_paragraph(" ".join(sentences))
    d.add_paragraph("We look forward to working with you.\n\nRavi Menon\nRegional Sales Director, PakAsia Global Pte Ltd")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# V4: angled phone photo of a printed rate card, priced per box (of 20)
# ---------------------------------------------------------------------------

def _perspective_coeffs(src, dst):
    matrix = []
    for (x, y), (X, Y) in zip(dst, src):
        matrix.append([x, y, 1, 0, 0, 0, -X * x, -X * y])
        matrix.append([0, 0, 0, x, y, 1, -Y * x, -Y * y])
    A = np.array(matrix, dtype=float)
    B = np.array([c for pt in src for c in pt], dtype=float)
    res = np.linalg.solve(A, B)
    return res.tolist()


def build_v4_photo(rfx: dict, prices: dict[int, float], rng: random.Random) -> bytes:
    W, H = 1500, 2100
    card = Image.new("RGB", (W, H), (252, 250, 245))
    dr = ImageDraw.Draw(card)
    f_title = ImageFont.load_default(size=46)
    f_head = ImageFont.load_default(size=30)
    f_row = ImageFont.load_default(size=27)
    f_small = ImageFont.load_default(size=19)
    dr.rectangle([0, 0, W, 130], fill=(120, 30, 30))
    dr.text((60, 35), "MEGHNA CORRUGATORS  -  RATE CARD 2026-27", font=f_title, fill=(255, 255, 255))
    dr.text((60, 95), "GIDC Vatva, Ahmedabad  |  Ph 98250 12345  |  meghnacorru@gmail.com", font=f_small, fill=(255, 230, 230))
    dr.text((60, 160), f"Party: Category Sourcing     Enquiry: {rfx['title'][:48]}", font=f_head, fill=(20, 20, 20))
    y = 220
    cols = [60, 140, 860, 1090, 1230]
    heads = ["Sr", "Item", "Size mm", "Ply", "Rate (Rs./Box)"]
    dr.rectangle([50, y - 10, W - 50, y + 42], fill=(225, 220, 210))
    for x, h in zip(cols, heads):
        dr.text((x, y), h, font=f_head, fill=(0, 0, 0))
    y += 56
    for li in rfx["line_items"]:
        rate_per_box = _round_up(prices[li["line_no"]] * 20, 5)  # box of 20 pieces
        desc = li["description"]
        if len(desc) > 40:
            desc = desc[:38] + ".."
        vals = [str(li["line_no"]), desc, f"{li['length_mm']}x{li['width_mm']}x{li['height_mm']}", li["board"].replace("-ply", "P"), f"{rate_per_box:,.0f}"]
        for x, v in zip(cols, vals):
            dr.text((x, y), v, font=f_row, fill=(15, 15, 15))
        dr.line([50, y + 44, W - 50, y + 44], fill=(190, 190, 190), width=2)
        y += 52
    dr.text((60, y + 20), "* All rates per box (bundle) of 20 nos., ex-factory Vatva. GST 12% extra. Payment 100% advance for first 3 orders.", font=f_small, fill=(60, 60, 60))
    dr.text((60, y + 48), "Rates valid 30 days. Transport extra. E&OE.", font=f_small, fill=(60, 60, 60))
    dr.text((W - 520, y + 90), "For Meghna Corrugators", font=f_small, fill=(60, 60, 60))
    dr.line([W - 520, y + 150, W - 140, y + 120], fill=(30, 30, 120), width=4)
    dr.line([W - 500, y + 135, W - 220, y + 150], fill=(30, 30, 120), width=3)

    # Photograph it: perspective, on a desk, slightly blurred, uneven light, bottom rows partly cut.
    out_w, out_h = 1300, 1750
    photo = Image.new("RGB", (out_w, out_h), (96, 74, 52))
    ImageDraw.Draw(photo)
    src = [(0, 0), (W, 0), (W, H), (0, H)]
    dst = [(150, 120), (1180, 60), (1270, 1660), (40, 1790)]  # bottom edge extends past the frame
    coeffs = _perspective_coeffs(src, dst)
    warped = card.transform((out_w, out_h), Image.PERSPECTIVE, coeffs, Image.BICUBIC, fillcolor=(96, 74, 52))
    mask = Image.new("L", (W, H), 255).transform((out_w, out_h), Image.PERSPECTIVE, coeffs, Image.BICUBIC, fillcolor=0)
    photo.paste(warped, (0, 0), mask)
    # lighting gradient
    grad = Image.new("L", (out_w, out_h))
    gd = ImageDraw.Draw(grad)
    for i in range(out_h):
        gd.line([(0, i), (out_w, i)], fill=int(235 - 70 * (i / out_h)))
    photo = Image.composite(photo, Image.new("RGB", (out_w, out_h), (0, 0, 0)), grad)
    photo = photo.rotate(-2.5, resample=Image.BICUBIC, fillcolor=(80, 62, 45))
    photo = photo.filter(ImageFilter.GaussianBlur(0.8))
    noise = np.array(photo).astype(np.int16)
    noise += rng.randint(-1, 1) + np.random.default_rng(rng.randint(0, 9999)).integers(-6, 7, size=noise.shape, dtype=np.int16)
    photo = Image.fromarray(np.clip(noise, 0, 255).astype(np.uint8))
    buf = io.BytesIO()
    photo.save(buf, format="JPEG", quality=78)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# V5: the one-line email
# ---------------------------------------------------------------------------

def build_v5_email(rfx: dict, prices: dict[int, float]) -> bytes:
    five = [li for li in rfx["line_items"] if li["board"] == "5-ply"]
    three = [li for li in rfx["line_items"] if li["board"] == "3-ply"]
    # Back out an implied per-kg rate from this vendor's per-piece prices and nominal weights.
    def implied_kg(items):
        if not items:
            return None
        return round(sum(prices[li["line_no"]] for li in items) / sum(li["nominal_weight_g"] / 1000 for li in items))

    kg5 = implied_kg(five) or 42
    kg3 = implied_kg(three) or 38
    body = (
        "From: Ganesh Board Mills <ganeshboard.indore@gmail.com>\n"
        "To: Category Sourcing\n"
        f"Date: {datetime.now().strftime('%a, %d %b %Y 18:42')} +0530\n"
        f"Subject: Re: RFQ: {rfx['title']}\n\n"
        f"Sir,\n\n₹{kg5}/kg for the 5-ply, {kg3} for the 3-ply, rest same as last year, freight extra.\n\n"
        "Regards\nDinesh Agrawal\nGanesh Board Mills, Sanwer Road, Indore\n(ISO 9001 certified unit)\n\n"
        "Sent from my phone\n"
    )
    return body.encode("utf-8")


# ---------------------------------------------------------------------------
# Certificates
# ---------------------------------------------------------------------------

def build_certificate_pdf(title: str, holder: str, body_lines: list[str]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.draw_rect(pymupdf.Rect(30, 30, 565, 812), color=(0.2, 0.3, 0.5), width=3)
    page.insert_text((60, 120), title, fontsize=22, fontname="hebo", color=(0.2, 0.3, 0.5))
    page.insert_text((60, 170), "This is to certify that", fontsize=11, fontname="helv")
    page.insert_text((60, 200), holder, fontsize=16, fontname="hebo")
    y = 240
    for line in body_lines:
        page.insert_text((60, y), line, fontsize=10.5, fontname="helv")
        y += 18
    page.insert_text((60, 700), "Certificate No. " + uuid.uuid4().hex[:12].upper(), fontsize=9, fontname="helv")
    out = doc.tobytes()
    doc.close()
    return out


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def _file_record(rfx_id: str, vendor_id: str, name: str, data: bytes, kind: str, content_type: str, role: str = "quote") -> dict:
    fid = uuid.uuid4().hex[:8]
    url = storage.put_bytes(f"rfx/{rfx_id}/files/{vendor_id}/{fid}-{name}", data, content_type)
    return {"file_id": fid, "name": name, "kind": kind, "content_type": content_type, "url": url, "size": len(data), "role": role}


def simulate_replies(state: dict) -> None:
    """Generate the five vendor replies and attach them to state['vendors'] (in place)."""
    rfx = state["rfx"]
    rfx_id = state["id"]
    seed = int(rfx_id[:6], 16)
    rng = random.Random(seed)
    prices = _vendor_prices(rfx, seed)
    usd_rate = state["fx"]["rates_to_inr"]["USD"] * 1.012  # vendor uses a slightly different rate than treasury
    if not state["vendors"]:
        state["vendors"] = [empty_vendor(v) for v in VENDORS]
    by_id = {v["vendor_id"]: v for v in state["vendors"]}
    received = datetime.now(timezone.utc)

    v = by_id["v1"]
    v["files"] = [
        _file_record(rfx_id, "v1", "SBP_Commercial_Offer_2026-27.xlsx", build_v1_xlsx(rfx, prices["v1"], rng), "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        _file_record(rfx_id, "v1", "SBP_ISO9001_Certificate.pdf", build_certificate_pdf("ISO 9001:2015 Certificate of Registration", "Sri Balaji Packaging Pvt Ltd, Chakan, Pune", ["operates a Quality Management System which complies with ISO 9001:2015", "Scope: Manufacture and supply of corrugated boxes and sheets.", "Valid from 12 March 2025 to 11 March 2028.", "Issued by TUV Nord Cert GmbH."]), "pdf", "application/pdf", "attachment"),
        _file_record(rfx_id, "v1", "SBP_FSC_CoC_Certificate.pdf", build_certificate_pdf("FSC Chain of Custody Certificate", "Sri Balaji Packaging Pvt Ltd", ["FSC-C0198765  |  License code FSC-C019876", "Scope: Purchase of FSC Mix kraft liner and manufacture of FSC Mix corrugated packaging.", "Valid to 30 June 2029."]), "pdf", "application/pdf", "attachment"),
    ]
    pdf_bytes, _skipped = build_v2_pdf(rfx, prices["v2"], rng)
    by_id["v2"]["files"] = [
        _file_record(rfx_id, "v2", "Kraftline_Quotation.pdf", pdf_bytes, "pdf", "application/pdf"),
        _file_record(rfx_id, "v2", "Kraftline_FSC_Certificate.pdf", build_certificate_pdf("FSC Chain of Custody Certificate", "Kraftline Industries, Nashik", ["FSC-C0245511", "Scope: Manufacture of corrugated packaging from FSC Mix and FSC Recycled paper.", "Valid to 14 November 2027."]), "pdf", "application/pdf", "attachment"),
    ]
    by_id["v3"]["files"] = [
        _file_record(rfx_id, "v3", "PakAsia_Commercial_Proposal.docx", build_v3_docx(rfx, prices["v3"], usd_rate, rng), "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ]
    by_id["v4"]["files"] = [
        _file_record(rfx_id, "v4", "IMG_20260918_174233.jpg", build_v4_photo(rfx, prices["v4"], rng), "image", "image/jpeg"),
    ]
    by_id["v5"]["files"] = [
        _file_record(rfx_id, "v5", "Re_RFQ_Ganesh_Board_Mills.eml", build_v5_email(rfx, prices["v5"]), "email", "message/rfc822"),
    ]
    for i, vid in enumerate(["v1", "v2", "v3", "v4", "v5"]):
        vv = by_id[vid]
        vv["status"] = "received"
        vv["extraction"] = None
        vv["texts"] = {}
        vv["error"] = None
        vv["received_at"] = (received + timedelta(days=i + 2, hours=i * 3)).isoformat(timespec="seconds")
    state["status"] = "responses_received"
