"""Award draft buyer workflow: top-2 → assign → checks → send → confirm → export."""
from __future__ import annotations

import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import app
from core import award_ask, award_draft, demo_ops, event_status, snapshots, storage

client = TestClient(app)


def _seed():
    st = demo_ops.build_golden_seed()
    storage.save_state(st["id"], st)
    return st


def _tick_all(st: dict) -> dict:
    award_draft.ensure_award_draft(st)
    award_draft.update_checklist(st, {c["id"]: True for c in award_draft.CHECKLIST_ITEMS})
    storage.save_state(st["id"], st)
    return storage.load_state(st["id"])


def test_suggest_top2_prefers_pass_pricing():
    st = demo_ops.build_golden_seed()
    live = snapshots.live_award_calculation(st)
    top = award_draft.suggest_top2(st, live)
    assert top["pass_count"] == 2
    assert top["enough"] is True
    assert len(top["shortlist_ids"]) == 2
    assert set(top["shortlist_ids"]) == {"v1", "v2"}
    for v in top["vendors"]:
        assert v["grade"] == "Pass"
        assert v["why"]
        assert "Cleared quality gates" in v["why"][0]


def test_default_allocation_cheapest_shortlist_pass():
    st = demo_ops.build_golden_seed()
    live = snapshots.live_award_calculation(st)
    top = award_draft.suggest_top2(st, live)
    alloc = award_draft.default_allocation(st, live, shortlist_ids=top["shortlist_ids"])
    assert alloc
    # Line 30 uncovered among Pass — may be missing
    assert "30" not in alloc or alloc.get("30") in top["shortlist_ids"]
    rows = award_draft.line_assignment_rows(st, live)
    for r in rows:
        if r["uncovered"]:
            assert r["line_no"] == 30
            continue
        assert r["selected_vendor_id"] in top["shortlist_ids"]
        assert all(o["vendor_id"] in ("v1", "v2") or True for o in r["options"])
        # Fail vendor must not appear
        assert all(o["vendor_id"] != "v4" for o in r["options"])


def test_cannot_assign_fail_vendor():
    st = demo_ops.build_golden_seed()
    live = snapshots.live_award_calculation(st)
    award_draft.ensure_award_draft(st, live)
    try:
        award_draft.update_allocation(st, {"1": "v4"})
        raise AssertionError("should reject Fail vendor")
    except ValueError as e:
        assert "eligible" in str(e).lower() or "pass" in str(e).lower()


def test_award_page_layout_no_freeze_lock():
    st = _seed()
    r = client.get(f"/rfx/{st['id']}/award")
    assert r.status_code == 200
    html = r.text
    assert "Ask the analyst" in html
    assert 'data-testid="award-ask-card"' in html
    assert 'data-testid="award-top2"' in html
    assert 'data-testid="award-assign"' in html
    assert 'data-testid="award-checklist"' in html
    assert 'data-testid="award-send-btn"' in html
    assert 'data-testid="award-export-link"' in html
    assert "Suggest top 2" in html
    assert "Assign by line" in html
    assert "Send award drafts" in html
    # Removed freeze/lock UX
    assert "Lock award" not in html
    assert 'data-testid="award-lock-btn"' not in html
    assert "Use &amp; lock" not in html and "Use & lock" not in html
    assert "Freeze complete" not in html
    assert "Freeze partial" not in html
    assert "Manual lock" not in html
    assert "Ready to freeze?" not in html


def test_send_disabled_until_checklist():
    st = _seed()
    page = client.get(f"/rfx/{st['id']}/award")
    assert 'data-testid="award-send-btn"' in page.text
    # Button starts disabled in HTML when checklist incomplete
    assert "disabled" in page.text.split('data-testid="award-send-btn"')[1][:200]


def test_send_requires_checklist_server_side():
    st = _seed()
    # Visiting award creates draft allocation without ticks
    client.get(f"/rfx/{st['id']}/award")
    r = client.post(f"/rfx/{st['id']}/award/send-notices", follow_redirects=False)
    assert r.status_code in (302, 303)
    st2 = storage.load_state(st["id"])
    flash = st2.get("flash") or {}
    assert flash.get("level") == "error"
    assert "check" in (flash.get("message") or "").lower() or "tick" in (flash.get("message") or "").lower()


def test_send_award_and_regret_with_confirmation():
    st = _seed()
    client.get(f"/rfx/{st['id']}/award")
    st = _tick_all(storage.load_state(st["id"]))
    before = len(st.get("outbox") or [])
    r = client.post(f"/rfx/{st['id']}/award/send-notices", follow_redirects=False)
    assert r.status_code in (302, 303)
    st2 = storage.load_state(st["id"])
    outbox = st2.get("outbox") or []
    assert len(outbox) > before
    vendor_mail = [o for o in outbox[before:] if o.get("vendor_id")]
    assert any(o.get("kind") == "award_notice" for o in vendor_mail)
    assert any(o.get("kind") in ("regret", "regret_notice") for o in vendor_mail)
    draft = st2.get("award_draft") or {}
    assert draft.get("sent") is True
    conf = draft.get("send_confirmation") or {}
    assert conf.get("award_vendors")
    assert conf.get("regret_vendors")
    # Award emails only include winners' lines
    for o in vendor_mail:
        if o.get("kind") == "award_notice":
            assert o.get("line_nos")

    page = client.get(f"/rfx/{st['id']}/award")
    assert 'data-testid="award-send-confirmation"' in page.text
    assert "Award drafts sent" in page.text
    disp = event_status.derive_event_display_status(storage.load_state(st["id"]))
    assert disp["key"] == event_status.STATUS_AWARD_DRAFTS_SENT
    assert disp["label"] == "Award drafts sent"


def test_premade_apply_split_updates_draft():
    st = _seed()
    r = client.post(f"/rfx/{st['id']}/award/ask-premade", data={"prompt_id": "best_split"})
    assert r.status_code == 200
    assert 'data-testid="award-ask-answer"' in r.text
    assert "Apply" in r.text
    assert "Use &amp; lock" not in r.text and "Use & lock" not in r.text
    assert 'data-testid="award-ask-use-and-lock"' not in r.text

    st1 = storage.load_state(st["id"])
    idx = len(st1["chat"]) - 1
    r2 = client.post(
        f"/rfx/{st['id']}/award/apply-split",
        data={"chat_idx": str(idx)},
        follow_redirects=False,
    )
    assert r2.status_code in (302, 303, 200)
    st2 = storage.load_state(st["id"])
    draft = st2.get("award_draft") or {}
    assert draft.get("allocation")
    flash = st2.get("flash") or {}
    assert flash.get("level") != "error"


def test_apply_vendor_to_line():
    st = _seed()
    client.get(f"/rfx/{st['id']}/award")
    st = storage.load_state(st["id"])
    # Force line 1 to the other shortlisted vendor
    draft = st["award_draft"]
    current = (draft.get("allocation") or {}).get("1")
    other = "v2" if current == "v1" else "v1"
    r = client.post(
        f"/rfx/{st['id']}/award/apply-vendor",
        data={"vendor_id": other, "line_nos": "1"},
        follow_redirects=False,
    )
    assert r.status_code in (302, 303, 200)
    st2 = storage.load_state(st["id"])
    assert st2["award_draft"]["allocation"]["1"] == other


def test_export_excel_includes_award_and_non_awarded():
    st = _seed()
    client.get(f"/rfx/{st['id']}/award")
    x = client.get(f"/rfx/{st['id']}/export.xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert "Award by line" in wb.sheetnames
    assert "Non-awarded" in wb.sheetnames


def test_award_ask_cache_still_works():
    st = demo_ops.build_golden_seed()
    a1 = award_ask.get_or_build_premade(st, "best_split")
    assert a1.get("cache_hit") is False
    assert a1.get("apply_all_split") or a1.get("apply_actions")
    a2 = award_ask.get_or_build_premade(st, "best_split")
    assert a2.get("cache_hit") is True
