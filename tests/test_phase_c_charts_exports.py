"""Phase C — charts helpers, export polish, audit strip."""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from core import charts, demo_ops, export, snapshots


def _seed():
    return demo_ops.build_golden_seed()


def test_chart_allocation_totals_match_live_share():
    st = _seed()
    live = snapshots.live_award_calculation(st)
    assert live["available"]
    bundle = charts.build_chart_bundle(st, live=live)
    assert bundle["available"]
    assert bundle["consistent"] is True

    alloc = bundle["allocation"]
    share = live["share_by_vendor"]
    assert alloc["total_lines"] == sum(int(s["lines"]) for s in share.values())
    assert abs(alloc["total_extended_inr"] - float(live["total_extended_inr"])) < 0.02
    assert abs(alloc["total_extended_inr"] - bundle["cost"]["total_extended_inr"]) < 0.02
    assert alloc["total_lines"] == bundle["coverage"]["covered"]
    assert bundle["coverage"]["covered"] + bundle["coverage"]["uncovered"] == bundle["coverage"]["total"]
    assert live["vendor_data_version"] == bundle["vendor_data_version"]
    assert live["snapshot"]["id"] == bundle["snapshot_id"]
    assert "v" in bundle["caption"] and "Snapshot" in bundle["caption"]


def test_chart_helpers_empty_without_extractions():
    st = _seed()
    for v in st["vendors"]:
        v["extraction"] = None
        v["status"] = "awaiting"
    bundle = charts.build_chart_bundle(st)
    assert bundle["available"] is False


def test_allocation_pcts_sum_near_100():
    share = {
        "A": {"lines": 10, "extended_inr": 750.0},
        "B": {"lines": 5, "extended_inr": 250.0},
    }
    alloc = charts.allocation_by_vendor(share)
    assert alloc["total_lines"] == 15
    assert alloc["total_extended_inr"] == 1000.0
    assert abs(sum(i["extended_pct"] for i in alloc["bars"]) - 100.0) < 0.2


def test_export_workbook_sheet_names_and_xlsx_bytes():
    st = _seed()
    data = export.award_workbook(st, provisional=False)
    assert data[:2] == b"PK"  # zip/xlsx
    wb = load_workbook(BytesIO(data))
    names = wb.sheetnames
    for required in ("Summary", "Award by line", "Comparison", "Flags", "Snapshot metadata"):
        assert required in names, names
    assert "Snapshot" not in names or "Snapshot metadata" in names
    assert "Snapshot" not in [n for n in names if n == "Snapshot"]


def test_export_includes_freeze_metadata_when_frozen():
    st = _seed()
    # Minimal freeze pack stamp
    live = snapshots.live_award_calculation(st)
    st.setdefault("freeze_packs", []).append(
        {
            "id": "fz-test-1",
            "status": "frozen",
            "freeze_mode": "partial",
            "strategy": "quality-gated cheapest per line",
            "calculation_snapshot_id": live["snapshot"]["id"],
            "vendor_data_version": live["vendor_data_version"],
            "total_extended_inr": live["total_extended_inr"],
            "frozen_at": "2026-09-20T00:00:00+00:00",
            "share_by_vendor": live["share_by_vendor"],
            "covered_line_count": live["covered_line_count"],
            "uncovered_lines": live["uncovered_lines"],
            "notices": [],
            "regrets": [],
        }
    )
    data = export.award_workbook(st)
    wb = load_workbook(BytesIO(data))
    ws = wb["Snapshot metadata"]
    fields = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert fields.get("Freeze ID") == "fz-test-1"
    assert fields.get("Freeze strategy") == "quality-gated cheapest per line"
    summary = wb["Summary"]["A2"].value or ""
    assert "fz-test-1" in summary
    assert "quality-gated" in summary

    memo = export.award_memo_md(st)
    assert "fz-test-1" in memo


def test_audit_trail_csv_and_trust_strip():
    st = _seed()
    st["reviews"] = [
        {
            "at": "2026-09-20T10:00:00+00:00",
            "vendor_name": "Kraftline Industries",
            "vendor_id": "kraftline",
            "line_no": 1,
            "action": "override",
            "value_inr": 12.5,
            "note": "test",
        }
    ]
    st["outbox"] = [
        {
            "kind": "award_notice",
            "vendor_name": "Kraftline Industries",
            "sent_at": "2026-09-20T11:00:00+00:00",
            "subject": "Award",
            "freeze_id": "fz-x",
        }
    ]
    csv_bytes = export.audit_trail_csv(st)
    text = csv_bytes.decode("utf-8")
    assert "section,at,vendor" in text
    assert "override" in text
    assert "award_notice" in text

    strip = charts.audit_trust_strip(st)
    assert strip["buyer_overrides"] == 1
    assert strip["notices_sent"] == 1
    assert "/anomalies" in strip["links"]["review"]
    assert "#outbox" in strip["links"]["outbox"]
    assert "/ai-log" in strip["links"]["ai_log"]


def test_provisional_export_still_gated_message():
    st = _seed()
    # Mark a vendor still processing
    st["vendors"][0]["status"] = "extracting"
    data = export.award_workbook(st, provisional=True)
    wb = load_workbook(BytesIO(data))
    a3 = wb["Summary"]["A3"].value or ""
    assert "PROVISIONAL" in a3
